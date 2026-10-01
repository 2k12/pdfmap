"""RF-14 (segundo plano/progreso/cancelación), RF-15 (descargas), RF-16 (estadísticas), RF-17 (filas paginadas)."""

import csv
import io

import pytest
from openpyxl import load_workbook

from tests.conftest import upload_file, wait_file_ready, wait_job
from tests.fixtures.generate_fixtures import build_report, write_pdf

API = "/api/v1"


@pytest.fixture
def done_job(client, ready_file, ventas_template):
    res = client.post(
        f"{API}/jobs", json={"file_id": ready_file["id"], "template": ventas_template, "formats": ["xlsx", "csv"]}
    )
    assert res.status_code == 202, res.text
    job = wait_job(client, res.json()["id"])
    assert job["status"] == "done", job
    return job


@pytest.mark.req("RF-14", "RF-16")
def test_job_completes_with_stats(done_job):
    assert done_job["progress"] == 1 and done_job["rows"] == 111
    assert done_job["stats"]["checks_ok"] == 15 and done_job["stats"]["checks_failed_count"] == 0
    assert done_job["formats"] == ["csv", "xlsx"]
    assert "111 filas" in done_job["message"]


@pytest.mark.req("RF-17")
def test_rows_pagination(client, done_job):
    page = client.get(f"{API}/jobs/{done_job['id']}/rows", params={"offset": 100, "limit": 50}).json()
    assert page["total"] == 111 and len(page["rows"]) == 11
    assert page["columns"][0] == {"id": "vendedor", "header": "Vendedor", "type": "text"}
    assert client.get(f"{API}/jobs/{done_job['id']}/rows", params={"limit": 5000}).status_code == 422


@pytest.mark.req("RF-15")
def test_download_xlsx_and_csv(client, done_job):
    res = client.get(f"{API}/jobs/{done_job['id']}/download", params={"format": "xlsx"})
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(res.content)).active
    assert ws.title == "Ventas" and ws["A1"].value == "Vendedor"
    res = client.get(f"{API}/jobs/{done_job['id']}/download", params={"format": "csv"})
    rows = list(csv.reader(io.StringIO(res.content.decode("utf-8-sig")), delimiter=";"))
    assert len(rows) == 112


@pytest.mark.req("RF-15")
def test_download_format_not_generated(client, ready_file, ventas_template):
    job = client.post(
        f"{API}/jobs", json={"file_id": ready_file["id"], "template": ventas_template, "formats": ["csv"]}
    ).json()
    wait_job(client, job["id"])
    assert client.get(f"{API}/jobs/{job['id']}/download", params={"format": "xlsx"}).status_code == 404
    assert client.get(f"{API}/jobs/{job['id']}/download", params={"format": "pdf"}).status_code == 422


@pytest.mark.req("RF-14", "RF-13")
def test_job_with_saved_template_id(client, ready_file, ventas_template):
    tpl = client.post(f"{API}/templates", json=ventas_template).json()
    job = client.post(f"{API}/jobs", json={"file_id": ready_file["id"], "template_id": tpl["id"]}).json()
    assert wait_job(client, job["id"])["rows"] == 111
    assert client.get(f"{API}/jobs", params={"file_id": ready_file["id"]}).json()[0]["id"] == job["id"]


@pytest.mark.req("RF-14")
def test_job_requires_exactly_one_template_source(client, ready_file, ventas_template):
    both = {"file_id": ready_file["id"], "template_id": "x", "template": ventas_template}
    assert client.post(f"{API}/jobs", json=both).status_code == 422
    assert client.post(f"{API}/jobs", json={"file_id": ready_file["id"]}).status_code == 422
    assert client.post(f"{API}/jobs", json={"file_id": ready_file["id"], "template_id": "nope"}).status_code == 404


@pytest.mark.req("RF-14")
def test_job_without_columns_rejected(client, ready_file, ventas_template):
    res = client.post(f"{API}/jobs", json={"file_id": ready_file["id"], "template": {**ventas_template, "columns": []}})
    assert res.status_code == 422


@pytest.mark.req("RF-14")
def test_cancel_job_and_download_conflict(client, tmp_path, ventas_template):
    pages, _ = build_report(pages=150)
    info = wait_file_ready(client, upload_file(client, write_pdf(pages, tmp_path / "grande.pdf"))["id"])
    job = client.post(f"{API}/jobs", json={"file_id": info["id"], "template": ventas_template}).json()
    client.post(f"{API}/jobs/{job['id']}/cancel")
    final = wait_job(client, job["id"])
    assert final["status"] in ("cancelled", "done")  # puede terminar antes de recibir la cancelación
    if final["status"] == "cancelled":
        assert client.get(f"{API}/jobs/{job['id']}/download").status_code == 409


@pytest.mark.req("RF-14")
def test_reextract_blocked_while_job_running(client, ready_file, ventas_template, monkeypatch):
    import threading

    from app.services import mapping

    gate = threading.Event()
    original = mapping.MappingService._counted

    def slow(self, ctx, file_id, total, engine):
        gate.wait(5)
        yield from original(self, ctx, file_id, total, engine)

    monkeypatch.setattr(mapping.MappingService, "_counted", slow)
    job = client.post(f"{API}/jobs", json={"file_id": ready_file["id"], "template": ventas_template}).json()
    assert client.post(f"{API}/files/{ready_file['id']}/extract", json={}).status_code == 409
    gate.set()
    assert wait_job(client, job["id"])["status"] == "done"


@pytest.mark.req("RF-14")
def test_unknown_job(client):
    assert client.get(f"{API}/jobs/00000000-0000-0000-0000-000000000000").status_code == 404
    assert client.post(f"{API}/jobs/no-es-uuid/cancel").status_code == 404
