"""Integración extractor PDF/TXT → motor → exportadores, con el reporte sintético."""

from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.core.export import ExcelExporter
from app.core.extraction import PdfExtractor, TextExtractor, iter_lines, open_extractor
from app.core.mapping import MappingEngine, column_types
from tests.conftest import engine_template
from tests.fixtures.generate_fixtures import build_report, write_pdf


@pytest.mark.req("RF-04")
def test_pdf_grid_reproduces_generated_lines(sample_pdf):
    pages, _ = build_report()
    with PdfExtractor(str(sample_pdf)) as ex:
        assert ex.page_count() == len(pages)
        assert ex.params.char_width == pytest.approx(4.2, abs=0.01)
        for n, expected in enumerate(pages, start=1):
            assert ex.page_lines(n) == [ln.rstrip() for ln in expected if ln.strip()]


@pytest.mark.req("RF-04")
def test_pdf_manual_params_override_calibration(sample_pdf):
    with PdfExtractor(str(sample_pdf), char_width=8.4, x_origin=20.0, y_tolerance=2.0) as ex:
        assert ex.params.char_width == 8.4
        assert ex.page_lines(1)[0].startswith("EMPRESA")
        with pytest.raises(IndexError):
            ex.page_lines(99)


@pytest.mark.req("RF-04")
def test_open_extractor_rejects_unknown_kind(sample_pdf):
    with pytest.raises(ValueError):
        open_extractor(sample_pdf, "docx")


@pytest.mark.req("RF-15", "RF-16", "RF-10")
def test_full_pipeline_to_excel(tmp_path, sample_pdf, ventas_template):
    tpl = engine_template(ventas_template)
    _, expected = build_report()
    engine = MappingEngine(tpl)
    out = tmp_path / "ventas.xlsx"
    writer = ExcelExporter(out, engine.columns, tpl["options"], column_types(tpl))
    with PdfExtractor(str(sample_pdf)) as ex:
        for row in engine.process(iter_lines(ex)):
            writer.write(row)
    writer.close()
    stats = engine.finish_stats()
    assert stats["rows"] == expected.details
    assert stats["checks_ok"] == expected.vendors and stats["checks_failed_count"] == 0
    assert stats["error_count"] == 0 and stats["unmatched"] == 0
    assert stats["rule_hits"]["nota"] == expected.notes

    ws = load_workbook(out)["Ventas"]
    assert ws.max_row == expected.details + 3  # encabezado + filas + blanco + totales
    total = sum(Decimal(str(ws.cell(r, 8).value)) for r in range(2, expected.details + 2))
    assert total == expected.total_amount
    notas = [ws.cell(r, 9).value for r in range(2, expected.details + 2) if ws.cell(r, 9).value]
    assert len(notas) == expected.notes and all(n.startswith("entrega parcial") for n in notas)


@pytest.mark.req("RF-16")
def test_pipeline_detects_corrupted_total(tmp_path, ventas_template):
    pages, expected = build_report(pages=2, corrupt_vendor=1)
    pdf = write_pdf(pages, tmp_path / "corrupto.pdf")
    engine = MappingEngine(engine_template(ventas_template))
    with PdfExtractor(str(pdf)) as ex:
        list(engine.process(iter_lines(ex)))
    stats = engine.finish_stats()
    assert stats["checks_failed_count"] == 1
    assert "total: impreso" in stats["checks_failed"][0]["message"]
    assert expected.corrupted == ["V002"]


@pytest.mark.req("RF-04", "RNF-07")
def test_pdf_and_txt_produce_identical_rows(sample_pdf, sample_txt, ventas_template):
    tpl = engine_template(ventas_template)

    def run(ex):
        rows = list(MappingEngine(tpl).process(iter_lines(ex)))
        ex.close()
        return [r[:-1] for r in rows]  # sin la página (el TXT pagina distinto)

    assert run(PdfExtractor(str(sample_pdf))) == run(TextExtractor(sample_txt))


@pytest.mark.req("RF-04")
def test_pdf_without_text_layer_is_detected(tmp_path):
    pdf = write_pdf([[]], tmp_path / "vacio.pdf")
    with PdfExtractor(str(pdf)) as ex:
        assert not ex.has_text()
        assert ex.page_lines(1) == []
