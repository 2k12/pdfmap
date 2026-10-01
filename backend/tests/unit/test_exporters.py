"""RF-15 · Exportación XLSX (formato, totales, múltiples hojas) y CSV."""

import csv
from datetime import datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.core.export import CsvExporter, ExcelExporter

pytestmark = pytest.mark.req("RF-15")

COLUMNS = [
    {"id": "nombre", "header": "Nombre"},
    {"id": "fecha", "header": "Fecha"},
    {"id": "monto", "header": "Monto", "total": True, "width": 14},
]
TYPES = {"nombre": "text", "fecha": "date", "monto": "number"}
ROWS = [
    ["Ana", datetime(2026, 1, 5), Decimal("10.50")],
    ["Luis", None, Decimal("-2.25")],
    ["Ñandú", datetime(2026, 2, 1), None],
]


def test_excel_contents_and_formats(tmp_path):
    path = tmp_path / "x.xlsx"
    ex = ExcelExporter(path, COLUMNS, {"sheet_name": "Hoja"}, TYPES)
    for row in ROWS:
        ex.write(row)
    ex.close()
    ws = load_workbook(path)["Hoja"]
    assert [c.value for c in ws[1]] == ["Nombre", "Fecha", "Monto"]
    assert ws["A2"].value == "Ana" and ws["C2"].value == 10.5
    assert ws["B2"].number_format == "dd/mm/yyyy"
    assert ws["C2"].number_format == "#,##0.00"
    assert ws["A6"].value == "TOTAL"
    assert ws["C6"].value == "=SUBTOTAL(9,C2:C4)"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:C4"
    assert ws.column_dimensions["C"].width == 14


def test_excel_without_totals_and_custom_format(tmp_path):
    path = tmp_path / "y.xlsx"
    cols = [{**COLUMNS[2], "number_format": "0.000"}]
    ex = ExcelExporter(path, cols, {"totals_row": False, "freeze_header": False, "autofilter": False}, TYPES)
    ex.write([Decimal("1.5")])
    ex.close()
    ws = load_workbook(path).active
    assert ws.max_row == 2
    assert ws["A2"].number_format == "0.000"
    assert ws.freeze_panes is None


def test_excel_rolls_over_to_new_sheets(tmp_path):
    path = tmp_path / "z.xlsx"
    ex = ExcelExporter(path, COLUMNS[:1], {"sheet_name": "Datos"}, TYPES, max_rows_per_sheet=6)  # 4 filas por hoja
    for i in range(10):
        ex.write([f"r{i}"])
    ex.close()
    wb = load_workbook(path)
    assert wb.sheetnames == ["Datos", "Datos (2)", "Datos (3)"]
    assert wb["Datos (3)"]["A2"].value == "r8"
    assert ex.total_rows == 10


def test_csv_export(tmp_path):
    path = tmp_path / "x.csv"
    ex = CsvExporter(path, COLUMNS, delimiter=";")
    for row in ROWS:
        ex.write(row)
    ex.close()
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM para Excel
    rows = list(csv.reader(path.open(encoding="utf-8-sig"), delimiter=";"))
    assert rows[0] == ["Nombre", "Fecha", "Monto"]
    assert rows[1] == ["Ana", "2026-01-05", "10.50"]
    assert rows[2] == ["Luis", "", "-2.25"]
    assert rows[3][0] == "Ñandú"


@pytest.mark.req("RNF-03")
def test_formula_injection_is_neutralized(tmp_path):
    payloads = ['=HYPERLINK("http://x")', "+1+1", "@SUM(A1)", "-abc"]
    xlsx = tmp_path / "inj.xlsx"
    ex = ExcelExporter(xlsx, COLUMNS[:1], {}, TYPES)
    for p in payloads:
        ex.write([p])
    ex.close()
    ws = load_workbook(xlsx).active
    assert all(ws.cell(r, 1).data_type == "s" for r in range(2, 2 + len(payloads)))

    path = tmp_path / "inj.csv"
    ex = CsvExporter(path, COLUMNS[:1])
    for p in [*payloads, "-12.50"]:
        ex.write([p])
    ex.close()
    values = [r[0] for r in csv.reader(path.open(encoding="utf-8-sig"), delimiter=";")][1:]
    assert values[:4] == ["'" + p for p in payloads]
    assert values[4] == "-12.50"  # los números negativos no se alteran
