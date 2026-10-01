"""Prueba de sistema E2E vía API: reproduce el flujo completo del usuario sobre un formato
DESCONOCIDO (sin plantilla previa), usando solo los asistentes de mapeo, igual que la UI:

subir por fragmentos → extraer → ver página → crear reglas desde líneas (regex sugerida) →
detectar columnas → anotar → diseñar columnas → vista previa → guardar plantilla →
procesar → resultados paginados → descargar Excel → eliminar.
"""

import io
import math
import re

import pytest
from openpyxl import load_workbook

from tests.conftest import wait_file_ready, wait_job
from tests.fixtures.generate_fixtures import build_report

API = "/api/v1"


@pytest.mark.req("RF-02", "RF-05", "RF-07", "RF-08", "RF-11", "RF-12", "RF-13", "RF-14", "RF-15", "RF-17", "RF-18")
def test_unknown_format_end_to_end(client, sample_pdf):
    # 1. Subida por fragmentos
    data = sample_pdf.read_bytes()
    up = client.post(f"{API}/uploads", json={"filename": "desconocido.pdf", "size": len(data)}).json()
    for i in range(math.ceil(len(data) / up["chunk_size"])):
        part = data[i * up["chunk_size"] : (i + 1) * up["chunk_size"]]
        assert client.put(f"{API}/uploads/{up['upload_id']}/chunks/{i}", content=part).status_code == 200
    file_id = client.post(f"{API}/uploads/{up['upload_id']}/complete").json()["id"]
    wait_file_ready(client, file_id)

    # 2. El usuario mira la página 1 y elige una línea de vendedor y una de detalle
    lines = [ln["text"] for ln in client.get(f"{API}/files/{file_id}/pages/1").json()["lines"]]
    vendor_line = next(ln for ln in lines if ln.startswith("Vendedor"))
    detail_line = next(ln for ln in lines if re.match(r"^\d{2}/", ln))

    # 3. Regex sugeridas a partir de esas líneas
    rx_vendor = client.post(f"{API}/mapping/suggest-regex", json={"line": vendor_line}).json()["regex"]
    rx_detail = client.post(f"{API}/mapping/suggest-regex", json={"line": detail_line, "tokens": 2}).json()["regex"]
    detail_lines = [ln for ln in lines if re.search(rx_detail, ln)]
    assert len(detail_lines) >= 5

    # 4. Detección automática de columnas en las líneas de detalle
    segments = client.post(f"{API}/mapping/suggest-columns", json={"lines": detail_lines, "min_gap": 2}).json()[
        "segments"
    ]
    fields = [
        {
            "key": f"c{i}",
            "start": s["start"],
            "end": s["end"],
            "type": s["type"],
            **({"date_formats": s["date_formats"]} if "date_formats" in s else {}),
            **({"decimal": s["decimal"]} if "decimal" in s else {}),
        }
        for i, s in enumerate(segments)
    ]
    template = {
        "name": "Descubierta en E2E",
        "rules": [
            {
                "id": "vendedor",
                "action": "context",
                "match": {"type": "regex", "value": rx_vendor},
                "fields": [{"key": "linea", "start": 10, "end": 14}],
            },
            {"id": "detalle", "action": "row", "match": {"type": "regex", "value": rx_detail}, "fields": fields},
        ],
        "columns": [{"id": "vendedor", "header": "Vendedor", "sources": ["vendedor.linea"]}]
        + [
            {
                "id": f["key"],
                "header": f"Col {i + 1}",
                "sources": [f"detalle.{f['key']}"],
                "total": f["type"] == "number",
            }
            for i, f in enumerate(fields)
        ],
    }

    # 5. Anotación y vista previa
    ann = client.post(f"{API}/mapping/annotate", json={"file_id": file_id, "template": template, "page": 1}).json()
    assert sum(ln["rule"] == "detalle" for ln in ann["lines"]) == len(detail_lines)
    preview = client.post(f"{API}/mapping/preview", json={"file_id": file_id, "template": template}).json()
    _, expected = build_report()
    assert preview["stats"]["rows"] == expected.details and preview["stats"]["error_count"] == 0
    assert preview["rows"][0][0] == "V001"

    # 6. Guardar plantilla y procesar
    saved = client.post(f"{API}/templates", json=template).json()
    job = client.post(
        f"{API}/jobs", json={"file_id": file_id, "template_id": saved["id"], "formats": ["xlsx", "csv"]}
    ).json()
    job = wait_job(client, job["id"])
    assert job["status"] == "done" and job["rows"] == expected.details

    # 7. Resultados paginados y descarga
    page = client.get(f"{API}/jobs/{job['id']}/rows", params={"limit": 10}).json()
    assert len(page["rows"]) == 10 and page["total"] == expected.details
    wb = load_workbook(io.BytesIO(client.get(f"{API}/jobs/{job['id']}/download").content))
    ws = wb.active
    assert ws.max_row == expected.details + 3
    number_cols = [i for i, f in enumerate(fields, start=2) if f["type"] == "number"]
    assert number_cols, "debería detectarse al menos una columna numérica (total)"
    total = sum(ws.cell(r, number_cols[-1]).value for r in range(2, expected.details + 2))
    assert round(total, 2) == float(expected.total_amount)

    # 8. Limpieza
    assert client.delete(f"{API}/files/{file_id}").status_code == 204
