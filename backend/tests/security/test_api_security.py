"""RNF-03 · Pruebas de seguridad (negativas) alineadas con OWASP API Security Top 10."""

import pytest
from fastapi.testclient import TestClient

from app.config import Settings
from app.main import create_app
from tests.conftest import upload_file, wait_file_ready

pytestmark = pytest.mark.req("RNF-03")
API = "/api/v1"


@pytest.fixture
def secured_client(tmp_path):
    settings = Settings(data_dir=tmp_path / "sec", api_key="s3cr3t-clave", chunk_size_mb=1)
    with TestClient(create_app(settings)) as c:
        yield c


def test_api_key_required_when_configured(secured_client):
    assert secured_client.get(f"{API}/files").status_code == 401
    assert secured_client.get(f"{API}/files", headers={"X-API-Key": "mala"}).status_code == 401
    assert secured_client.get(f"{API}/files", headers={"X-API-Key": "s3cr3t-clave"}).status_code == 200
    assert secured_client.get(f"{API}/health").status_code == 200  # la sonda de salud es pública


@pytest.mark.parametrize(
    "path",
    [
        "/files/..%2F..%2Fetc%2Fpasswd",
        "/files/../../etc/passwd/pages/1",
        "/uploads/..%5C..%5Cwindows",
        "/jobs/..%2Fpdfmap.sqlite3/download",
        "/files/1%20OR%201=1",
    ],
)
def test_path_traversal_and_injection_in_ids(client, path):
    assert client.get(API + path).status_code in (404, 422)


def test_uploaded_filename_is_sanitized(client, sample_pdf, settings):
    info = upload_file(client, sample_pdf, name="../../../evil<script>.pdf")
    assert info["name"] == "evil_script_.pdf"
    stored = list(settings.files_dir.rglob("*"))
    assert all(settings.files_dir in p.parents for p in stored)
    assert not any("evil" in p.name for p in stored)


def test_chunk_larger_than_declared_is_rejected(client, settings):
    up = client.post(f"{API}/uploads", json={"filename": "a.pdf", "size": 10}).json()
    res = client.put(f"{API}/uploads/{up['upload_id']}/chunks/0", content=b"x" * (settings.chunk_size + 1))
    assert res.status_code == 413


def test_regex_length_limits(client, ready_file, ventas_template):
    long_pattern = "a" * 1001
    res = client.get(f"{API}/files/{ready_file['id']}/search", params={"pattern": long_pattern, "mode": "regex"})
    assert res.status_code == 422
    bad = {**ventas_template}
    bad["rules"] = [{**ventas_template["rules"][0], "match": {"type": "regex", "value": "a" * 2000}}]
    assert client.post(f"{API}/templates", json=bad).status_code == 422


def test_sql_like_input_is_inert(client, ready_file):
    res = client.get(f"{API}/files/{ready_file['id']}/search", params={"pattern": "'; DROP TABLE lines; --"})
    assert res.status_code == 200 and res.json()["matches"] == []
    assert client.get(f"{API}/files/{ready_file['id']}/pages/1").status_code == 200


def test_security_headers(client):
    res = client.get(f"{API}/health")
    assert res.headers["x-content-type-options"] == "nosniff"
    assert res.headers["x-frame-options"] == "DENY"
    assert res.headers["referrer-policy"] == "no-referrer"


def test_cors_only_allows_configured_origins(client):
    ok = client.options(
        f"{API}/files", headers={"Origin": "http://localhost:5173", "Access-Control-Request-Method": "GET"}
    )
    assert ok.headers.get("access-control-allow-origin") == "http://localhost:5173"
    bad = client.options(
        f"{API}/files", headers={"Origin": "https://evil.example", "Access-Control-Request-Method": "GET"}
    )
    assert "access-control-allow-origin" not in bad.headers


def test_mass_assignment_is_prevented(client, ventas_template):
    res = client.post(f"{API}/templates", json={**ventas_template, "builtin": True, "id": "builtin-ventas-ejemplo"})
    assert res.status_code == 201
    created = res.json()
    assert created["builtin"] is False and created["id"] != "builtin-ventas-ejemplo"
    assert client.post(f"{API}/templates", json={**ventas_template, "rol": "admin"}).status_code == 422


def test_excel_formula_injection_is_not_evaluated(client, tmp_path, ventas_template):
    """Texto que empieza con '=' en el PDF debe llegar como texto, no como fórmula ejecutable."""
    from openpyxl import load_workbook

    from tests.fixtures.generate_fixtures import write_txt

    cliente = "=CMD|' /C calc'!A0"
    detalle = f"{'05/07/2026':<10}  {'F-000101':<8}  {cliente:<28}{1:>6}{'1.00':>12}"
    txt = write_txt([['Vendedor: V001 =HYPERLINK("http://x")     Zona: NORTE', detalle]], tmp_path / "inj.txt")
    info = wait_file_ready(client, upload_file(client, txt)["id"])
    job = client.post(f"{API}/jobs", json={"file_id": info["id"], "template": ventas_template}).json()
    from tests.conftest import wait_job

    assert wait_job(client, job["id"])["status"] == "done"
    content = client.get(f"{API}/jobs/{job['id']}/download").content
    out = tmp_path / "inj.xlsx"
    out.write_bytes(content)
    ws = load_workbook(out).active
    for cell in (ws["B2"], ws["F2"]):
        assert cell.data_type == "s", f"{cell.coordinate} se interpretó como fórmula: {cell.value}"
